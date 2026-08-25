from __future__ import annotations

from io import BytesIO
from pathlib import Path
from urllib.request import Request, urlopen

from PIL import Image, ImageDraw, ImageFilter
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT_ROOT / "entregables-cliente" / "Propuesta_Comercial_MaryTere_Aimarktech.docx"

# Identidad Aimarktech vigente
BLUE = "0A74DA"
CYAN = "00C2FF"
YELLOW = "FFCE00"
GREEN = "28A745"
GRAY = "D3D3D3"
INK = "0E1B2C"
INK_SOFT = "46566B"
BG_SOFT = "F4F8FD"
DARK = "071426"
WHITE = "FFFFFF"
LINE = "E3EBF4"
CREAM = "FFFDF2"
MARY_GOLD = "F2C14E"
MARY_TURQ = "5FB2C4"
MARY_SAGE = "A8C3A0"

PHONE = "+52 56 3963 7740"
EMAIL = "contacto@soyaimarktech.com"
WEBSITE = "https://soyaimarktech.com"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")


def set_repeat_header(row) -> None:
    set_repeat_table_header(row)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.name = "Lato"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("93A3B6")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run(run, *, size=None, bold=None, color=None, font="Lato", italic=None) -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def style_paragraph(paragraph, *, before=0, after=6, line=1.12, keep=False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if keep:
        fmt.keep_with_next = True


def add_text(doc_or_cell, text: str, *, size=9.4, color=INK, bold=False,
             italic=False, align=None, before=0, after=6, line=1.18) -> object:
    p = doc_or_cell.add_paragraph()
    if align is not None:
        p.alignment = align
    style_paragraph(p, before=before, after=after, line=line)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_rich_paragraph(doc_or_cell, parts, *, size=9.4, align=None,
                       before=0, after=6, line=1.18) -> object:
    p = doc_or_cell.add_paragraph()
    if align is not None:
        p.alignment = align
    style_paragraph(p, before=before, after=after, line=line)
    for text, bold, color, italic in parts:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, color=color or INK, italic=italic)
    return p


def add_bullet(doc_or_cell, text: str, *, level=0, color=INK, size=9.1,
               bold_prefix: str | None = None) -> object:
    p = doc_or_cell.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    style_paragraph(p, after=3.5, line=1.12)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run(r1, size=size, bold=True, color=color)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run(r2, size=size, color=color)
    else:
        r = p.add_run(text)
        set_run(r, size=size, color=color)
    return p


def add_numbered(doc_or_cell, text: str, level=0) -> object:
    p = doc_or_cell.add_paragraph(style="List Number" if level == 0 else "List Number 2")
    style_paragraph(p, after=4, line=1.13)
    r = p.add_run(text)
    set_run(r, size=9.2, color=INK)
    return p


def add_section_heading(doc, number: str, title: str, subtitle: str | None = None,
                        color=BLUE) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, before=8, after=5, keep=True)
    r = p.add_run(f"{number}  ")
    set_run(r, size=10, bold=True, color=color, font="Montserrat")
    r = p.add_run(title)
    set_run(r, size=17, bold=True, color=INK, font="Montserrat")
    # Línea de acento
    p2 = doc.add_paragraph()
    style_paragraph(p2, after=8, keep=True)
    p2.paragraph_format.space_before = Pt(0)
    r2 = p2.add_run("━━━━━━━━━━━━")
    set_run(r2, size=6, bold=True, color=color)
    if subtitle:
        p3 = doc.add_paragraph()
        style_paragraph(p3, after=10)
        r3 = p3.add_run(subtitle)
        set_run(r3, size=9.1, italic=True, color=INK_SOFT)


def add_callout(doc_or_cell, title: str, body: str, *, fill=BG_SOFT,
                stripe=BLUE, title_color=INK, body_color=INK_SOFT) -> object:
    table = doc_or_cell.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_layout(table)
    table.columns[0].width = Cm(0.25)
    table.columns[1].width = Cm(16.5)
    set_cell_shading(table.cell(0, 0), stripe)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_margins(table.cell(0, 0), 0, 0, 0, 0)
    set_cell_margins(table.cell(0, 1), 160, 220, 160, 220)
    table.cell(0, 0).text = ""
    c = table.cell(0, 1)
    c.text = ""
    p = c.paragraphs[0]
    style_paragraph(p, after=4, keep=True)
    r = p.add_run(title)
    set_run(r, size=10, bold=True, color=title_color, font="Montserrat")
    p2 = c.add_paragraph()
    style_paragraph(p2, after=0, line=1.15)
    r2 = p2.add_run(body)
    set_run(r2, size=9.1, color=body_color)
    for cell in table.rows[0].cells:
        set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
    return table


def add_spacer(doc, pts=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)
    p.paragraph_format.space_before = Pt(0)
    p.add_run("")


def add_table(doc, headers, rows, widths=None, header_fill=DARK,
              alternating=True, font_size=8.5, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_layout(table)
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Cm(width)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, 120, 120, 120, 120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_paragraph(p, after=0)
        r = p.add_run(text)
        set_run(r, size=8.3, bold=True, color=WHITE, font="Montserrat")
        set_cell_border(cell, bottom={"val": "single", "sz": 6, "color": WHITE},
                        top={"val": "single", "sz": 6, "color": header_fill},
                        left={"val": "single", "sz": 3, "color": LINE},
                        right={"val": "single", "sz": 3, "color": LINE})
    for row_idx, row_data in enumerate(rows):
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        fill = WHITE if (not alternating or row_idx % 2 == 0) else BG_SOFT
        for col_idx, value in enumerate(row_data):
            cell = cells[col_idx]
            set_cell_shading(cell, fill)
            set_cell_margins(cell, 105, 120, 105, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line=1.08)
            r = p.add_run(str(value))
            set_run(r, size=font_size, bold=(first_col_bold and col_idx == 0), color=INK)
            set_cell_border(cell, bottom={"val": "single", "sz": 4, "color": LINE},
                            top={"val": "single", "sz": 4, "color": LINE},
                            left={"val": "single", "sz": 4, "color": LINE},
                            right={"val": "single", "sz": 4, "color": LINE})
    return table


def make_marytere_motif() -> BytesIO:
    """Motivo editorial inspirado en sol/mar; no reproduce el logotipo adjunto."""
    w, h = 1600, 520
    base = Image.new("RGBA", (w, h), (255, 255, 255, 0))
    glow = Image.new("RGBA", (w, h), (255, 255, 255, 0))
    gd = ImageDraw.Draw(glow, "RGBA")
    center = (1180, 150)
    for radius, alpha in [(220, 18), (170, 28), (120, 45), (75, 100)]:
        gd.ellipse((center[0]-radius, center[1]-radius, center[0]+radius, center[1]+radius), fill=(242, 193, 78, alpha))
    glow = glow.filter(ImageFilter.GaussianBlur(18))
    base.alpha_composite(glow)
    draw = ImageDraw.Draw(base, "RGBA")
    # Rayos sutiles
    for i in range(30):
        x = 920 + i * 20
        draw.line((center[0], center[1], x, 0 if i % 2 == 0 else 30), fill=(242, 193, 78, 32), width=3)
    # Olas abstractas
    for y, color, alpha in [(300, (95, 178, 196), 95), (340, (0, 194, 255), 65), (390, (10, 116, 218), 42)]:
        pts = []
        for x in range(-40, w + 80, 40):
            yy = y + (16 if (x // 40) % 2 == 0 else -10)
            pts.append((x, yy))
        pts += [(w, h), (0, h)]
        draw.polygon(pts, fill=(*color, alpha))
    # Hojas/puntos florales mínimos
    for x, y in [(70,410),(120,390),(170,430),(230,400),(290,445),(350,415),(430,450),(500,425)]:
        draw.ellipse((x-14,y-8,x+14,y+8), fill=(168,195,160,110))
        draw.ellipse((x+8,y-18,x+26,y), fill=(242,193,78,80))
    out = BytesIO()
    base.save(out, format="PNG")
    out.seek(0)
    return out


def fetch_aimarktech_logo() -> BytesIO | None:
    try:
        req = Request("https://soyaimarktech.com/assets/logo.png", headers={"User-Agent": "Mozilla/5.0"})
        data = urlopen(req, timeout=15).read()
        bio = BytesIO(data)
        Image.open(bio).verify()
        bio.seek(0)
        return bio
    except Exception:
        return None


def configure_document(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.55)
    sec.bottom_margin = Cm(1.55)
    sec.left_margin = Cm(1.75)
    sec.right_margin = Cm(1.75)
    sec.header_distance = Cm(0.65)
    sec.footer_distance = Cm(0.65)
    sec.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Lato"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Lato")
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name in ["List Bullet", "List Bullet 2", "List Number", "List Number 2"]:
        st = styles[style_name]
        st.font.name = "Lato"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Lato")
        st.font.size = Pt(9.1)
        st.font.color.rgb = RGBColor.from_string(INK)

    # Header: marca discreta
    header = sec.header
    table = header.add_table(rows=1, cols=2, width=Cm(17.4))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Cm(10.5)
    table.columns[1].width = Cm(6.9)
    c1, c2 = table.rows[0].cells
    c1.text = ""
    p = c1.paragraphs[0]
    r = p.add_run("AIMARKTECH")
    set_run(r, size=8.5, bold=True, color=BLUE, font="Montserrat")
    r = p.add_run("  ·  Propuesta confidencial")
    set_run(r, size=7.5, color="7F93A6")
    c2.text = ""
    p = c2.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Psicól. MaryTere")
    set_run(r, size=7.8, color=INK_SOFT)
    for c in (c1, c2):
        set_cell_margins(c, 0, 0, 40, 0)
        set_cell_border(c, bottom={"val": "single", "sz": 8, "color": CYAN})

    # Footer
    footer = sec.footer
    table = footer.add_table(rows=1, cols=2, width=Cm(17.4))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Cm(13.5)
    table.columns[1].width = Cm(3.9)
    c1, c2 = table.rows[0].cells
    c1.text = ""
    p = c1.paragraphs[0]
    r = p.add_run(f"{WEBSITE.replace('https://', '')}  ·  {PHONE}")
    set_run(r, size=7.5, color="7F93A6")
    c2.text = ""
    add_page_number(c2.paragraphs[0])
    for c in (c1, c2):
        set_cell_margins(c, 45, 0, 0, 0)
        set_cell_border(c, top={"val": "single", "sz": 6, "color": LINE})


def add_cover(doc: Document) -> None:
    # Marca Aimarktech
    logo = fetch_aimarktech_logo()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(p, after=10)
    if logo:
        try:
            p.add_run().add_picture(logo, width=Inches(1.75))
        except Exception:
            r = p.add_run("AIMARKTECH")
            set_run(r, size=18, bold=True, color=BLUE, font="Montserrat")
    else:
        r = p.add_run("AIMARKTECH")
        set_run(r, size=18, bold=True, color=BLUE, font="Montserrat")

    p = doc.add_paragraph()
    style_paragraph(p, after=3)
    r = p.add_run("PROPUESTA ESTRATÉGICA Y COMERCIAL")
    set_run(r, size=10.5, bold=True, color=BLUE, font="Montserrat")

    p = doc.add_paragraph()
    style_paragraph(p, after=7, line=1.0)
    r = p.add_run("Presencia digital y\ncaptación ética de pacientes")
    set_run(r, size=27, bold=True, color=INK, font="Montserrat")

    p = doc.add_paragraph()
    style_paragraph(p, after=14)
    r = p.add_run("Facebook · Instagram · WhatsApp · Meta Ads")
    set_run(r, size=11.5, bold=True, color=GREEN, font="Montserrat")

    doc.add_picture(make_marytere_motif(), width=Inches(6.55), height=Inches(2.12))

    add_spacer(doc, 3)
    p = doc.add_paragraph()
    style_paragraph(p, after=3)
    r = p.add_run("PREPARADA PARA")
    set_run(r, size=8, bold=True, color="7F93A6", font="Montserrat")
    p = doc.add_paragraph()
    style_paragraph(p, after=2)
    r = p.add_run("Psicól. MaryTere")
    set_run(r, size=18, bold=True, color=MARY_TURQ, font="Montserrat")
    p = doc.add_paragraph()
    style_paragraph(p, after=11)
    r = p.add_run("Especialista en Psicoterapia Gestalt · Querétaro y atención en línea")
    set_run(r, size=9.5, color=INK_SOFT)

    info = doc.add_table(rows=3, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_fixed_layout(info)
    info.columns[0].width = Cm(4.0)
    info.columns[1].width = Cm(12.6)
    details = [
        ("Elaborada por", "José Antonio Aguilar · Aimarktech"),
        ("Fecha", "11 de agosto de 2026"),
        ("Vigencia", "30 días naturales"),
    ]
    for row, (label, value) in zip(info.rows, details):
        set_row_cant_split(row)
        for i, text in enumerate((label, value)):
            cell = row.cells[i]
            set_cell_shading(cell, BG_SOFT if i == 0 else WHITE)
            set_cell_margins(cell, 100, 120, 100, 120)
            p = cell.paragraphs[0]
            style_paragraph(p, after=0)
            r = p.add_run(text)
            set_run(r, size=8.3, bold=(i == 0), color=INK if i else BLUE)
            set_cell_border(cell, bottom={"val": "single", "sz": 4, "color": LINE})

    add_spacer(doc, 6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=0)
    r = p.add_run(f"{PHONE}   ·   {EMAIL}   ·   {WEBSITE.replace('https://', '')}")
    set_run(r, size=8.3, bold=True, color=INK_SOFT)
    doc.add_page_break()


def add_executive_summary(doc: Document) -> None:
    add_section_heading(doc, "01", "Resumen ejecutivo",
                        "Una estrategia sencilla de operar, profesional por dentro y coherente con la ética terapéutica.")
    add_rich_paragraph(doc, [
        ("Estimada MaryTere:\n\n", True, INK, False),
        ("Gracias por compartirnos tu visión. Esta propuesta parte de una idea muy clara: ", False, INK, False),
        ("darte a conocer sin convertir la terapia en una venta agresiva", True, BLUE, False),
        (". Construiremos una presencia digital cálida y confiable que invite a las personas a conocerte, comprender tu enfoque y dar el paso de solicitar una cita.", False, INK, False),
    ], size=9.8, after=10, line=1.22)

    add_callout(doc, "La oportunidad central",
                "MaryTere ya cuenta con experiencia, certificación, cédula profesional y una filosofía auténtica. El reto no es inventar una promesa: es volver visible esa confianza y crear un camino sencillo desde una publicación hasta una conversación por WhatsApp.",
                fill=CREAM, stripe=YELLOW)
    add_spacer(doc, 5)

    add_text(doc, "Lo que entendimos de tu práctica", size=12, bold=True, color=INK, before=3, after=7)
    add_table(doc,
              ["Fortaleza actual", "Implicación estratégica"],
              [
                  ("4 años de experiencia, certificación y cédula", "La comunicación puede apoyarse en credenciales reales, no en promesas."),
                  ("Atención presencial y en línea", "Podemos captar localmente en Querétaro y ampliar cobertura digital a México."),
                  ("Estilo honesto: “invitar, no vender”", "El contenido será educativo, cálido y sin presión comercial."),
                  ("No deseas mostrar tu rostro", "Crearemos una marca faceless: ilustración, voz, palabras, consultorio y recursos visuales."),
                  ("Facebook nuevo y sin estrategia previa", "Podemos construir una base ordenada desde el inicio, sin corregir vicios anteriores."),
              ],
              widths=[6.3, 10.4], font_size=8.5, first_col_bold=True)

    add_spacer(doc, 7)
    add_callout(doc, "Resultado que buscamos en 90 días",
                "Una identidad coherente, perfiles profesionales, contenido constante, campañas controladas y un flujo medible: contenido → perfil → WhatsApp → cita. No prometemos un número artificial de pacientes; sí un sistema que permita aprender, medir y mejorar.",
                fill=BG_SOFT, stripe=GREEN)


def add_diagnosis(doc: Document) -> None:
    add_section_heading(doc, "02", "Diagnóstico y enfoque",
                        "El principal cuello de botella está arriba del embudo: hoy pocas personas conocen a MaryTere.")
    add_text(doc, "Situación actual", size=12, bold=True, color=INK, after=6)
    for text in [
        "Presencia digital inicial: Facebook recién abierto; Instagram y otros canales todavía sin desarrollar.",
        "Público descrito como “mayores de 18 años”, demasiado amplio para que el mensaje conecte con fuerza.",
        "WhatsApp Business disponible, pero falta definir el número dedicado y el proceso de respuesta/agendamiento.",
        "Logo propio con una estética cálida y natural, pero sin manual de marca ni formatos preparados para redes.",
        "Un testimonio disponible, sujeto a consentimiento expreso y protección de identidad.",
    ]:
        add_bullet(doc, text)

    add_spacer(doc, 5)
    add_text(doc, "Nicho de comunicación propuesto", size=12, bold=True, color=INK, after=6)
    add_callout(doc, "“Conocerte y aceptarte”",
                "La cuenta hablará principalmente a personas adultas que atraviesan preguntas personales, miedos, conflictos familiares o momentos de cambio y que buscan un espacio honesto para comprenderse. Este enfoque orienta el mensaje sin excluir otros motivos de consulta.",
                fill="EFFAFC", stripe=MARY_TURQ)

    add_spacer(doc, 7)
    add_text(doc, "Posicionamiento", size=12, bold=True, color=INK, after=5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, before=2, after=9, line=1.18)
    r = p.add_run("“Acompañamiento Gestalt honesto y cálido, sin fórmulas mágicas, para quienes desean conocerse y aceptarse.”")
    set_run(r, size=13, bold=True, italic=True, color=BLUE, font="Montserrat")

    add_text(doc, "Principios de comunicación", size=12, bold=True, color=INK, after=6)
    add_table(doc,
              ["Sí haremos", "No haremos"],
              [
                  ("Educar, acompañar e invitar con calidez", "Presionar, alarmar o usar miedo para vender"),
                  ("Comunicar credenciales y límites con claridad", "Prometer sanación, resultados o plazos clínicos"),
                  ("Hablar de emociones y momentos de vida", "Afirmar o inferir diagnósticos sobre la audiencia"),
                  ("Usar testimonios solo con permiso escrito", "Exponer pacientes, historias o datos sensibles"),
              ],
              widths=[8.35, 8.35], header_fill=INK, font_size=8.5)


def add_strategy(doc: Document) -> None:
    add_section_heading(doc, "03", "Estrategia propuesta",
                        "El rostro de la marca será su universo visual, su voz y sus palabras.")
    add_text(doc, "3.1 Marca faceless con humanidad", size=12, bold=True, color=INK, after=5)
    add_text(doc,
             "No mostrar el rostro no significa construir una marca fría. La estrategia combinará el lenguaje visual del sol, el agua y la naturaleza con recursos que transmitan presencia humana sin exponer a MaryTere.",
             after=7)
    for text in [
        "Identidad ilustrada consistente: colores, tipografías y plantillas reconocibles.",
        "Voz en off opcional para Reels y reflexiones; puede incorporarse gradualmente.",
        "Escenas sin rostro: manos, cuaderno, plantas, consultorio, luz, objetos y ejercicios escritos.",
        "Credenciales, metodología y límites explicados en lenguaje sencillo.",
        "Testimonio anónimo únicamente después de obtener consentimiento escrito.",
    ]:
        add_bullet(doc, text)

    add_spacer(doc, 5)
    add_text(doc, "3.2 Pilares de contenido", size=12, bold=True, color=INK, after=6)
    add_table(doc,
              ["Pilar", "Peso", "Propósito", "Ejemplos"],
              [
                  ("Educar", "40%", "Demostrar criterio profesional", "Gestalt en palabras simples; mitos; ejercicios breves"),
                  ("Conectar", "30%", "Generar identificación y guardados", "Preguntas de reflexión; autoaceptación; vínculos"),
                  ("Generar confianza", "20%", "Reducir dudas antes de escribir", "Cómo es una sesión; credenciales; enfoque honesto"),
                  ("Invitar", "10%", "Facilitar el siguiente paso", "Modalidades, horarios, cupos y enlace a WhatsApp"),
              ],
              widths=[3.2, 1.5, 5.0, 7.0], font_size=8.1, first_col_bold=True)

    add_spacer(doc, 7)
    add_text(doc, "3.3 Embudo sencillo", size=12, bold=True, color=INK, after=6)
    flow = doc.add_table(rows=1, cols=4)
    flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_layout(flow)
    steps = [
        ("1", "CONTENIDO", "Educa y conecta"),
        ("2", "PERFIL", "Explica y genera confianza"),
        ("3", "WHATSAPP", "Resuelve dudas y agenda"),
        ("4", "CITA", "Presencial o en línea"),
    ]
    colors = [BLUE, CYAN, GREEN, YELLOW]
    for i, (n, title, sub) in enumerate(steps):
        c = flow.cell(0, i)
        set_cell_shading(c, BG_SOFT if i != 3 else CREAM)
        set_cell_margins(c, 170, 100, 170, 100)
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=3)
        r = p.add_run(n)
        set_run(r, size=16, bold=True, color=colors[i], font="Montserrat")
        p2 = c.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p2, after=3)
        r2 = p2.add_run(title)
        set_run(r2, size=8.2, bold=True, color=INK, font="Montserrat")
        p3 = c.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p3, after=0, line=1.05)
        r3 = p3.add_run(sub)
        set_run(r3, size=7.3, color=INK_SOFT)
        set_cell_border(c, bottom={"val": "single", "sz": 8, "color": colors[i]},
                        top={"val": "single", "sz": 3, "color": LINE},
                        left={"val": "single", "sz": 3, "color": LINE},
                        right={"val": "single", "sz": 3, "color": LINE})

    add_spacer(doc, 8)
    add_callout(doc, "Canales prioritarios",
                "Arranque: Instagram + Facebook + WhatsApp Business. Fase posterior: Perfil de Empresa en Google y TikTok, solo cuando el sistema inicial ya sea sostenible. Abrir más canales no ayuda si no podemos atenderlos con constancia.",
                fill=BG_SOFT, stripe=BLUE)


def add_roadmap(doc: Document) -> None:
    add_section_heading(doc, "04", "Plan de trabajo · 90 días",
                        "Primero construimos la base; después publicamos, amplificamos y optimizamos.")
    add_table(doc,
              ["Fase", "Periodo", "Acciones principales", "Resultado"],
              [
                  ("0 · Cimientos", "Semanas 1–2", "Adaptación visual; mini guía; perfiles; WhatsApp; mensajes y accesos", "Canales listos y coherentes"),
                  ("1 · Siembra", "Semanas 3–5", "Calendario; lote inicial; publicaciones; stories; primeras conversaciones", "Presencia constante y primeras señales"),
                  ("2 · Pauta", "Semanas 6–9", "Campaña de alcance/interacción; campaña de mensajes; pruebas controladas", "Mayor visibilidad y contactos medibles"),
                  ("3 · Ajuste", "Semanas 10–12", "Análisis; optimización; reporte; plan del siguiente trimestre", "Sistema basado en aprendizajes reales"),
              ],
              widths=[2.7, 2.5, 7.3, 4.2], font_size=8.2, first_col_bold=True)

    add_spacer(doc, 8)
    add_text(doc, "Entregables incluidos en la configuración inicial", size=12, bold=True, color=INK, after=6)
    for text in [
        "Adaptación digital del logo existente para perfil, portada y publicaciones (no incluye rediseño o vectorización completa).",
        "Mini guía visual: colores, tipografías, tono y reglas básicas de aplicación.",
        "Kit de 6 plantillas editables en Canva.",
        "Creación u optimización de Instagram, Facebook y WhatsApp Business.",
        "Bio, descripción de servicios, mensaje de bienvenida y respuestas rápidas.",
        "Calendario editorial del primer mes y primer lote de contenido.",
        "Recurso descargable inicial: “5 preguntas para empezar a conocerte”.",
        "Configuración inicial de Meta Business Suite y estructura de campañas.",
    ]:
        add_bullet(doc, text)

    add_spacer(doc, 7)
    add_text(doc, "Ejemplos del primer mes", size=12, bold=True, color=INK, after=6)
    add_table(doc,
              ["Formato", "Tema", "Objetivo"],
              [
                  ("Carrusel", "¿Qué es la terapia Gestalt en palabras simples?", "Educar"),
                  ("Reel faceless", "3 señales de que te exiges de más", "Alcance"),
                  ("Post de marca", "Conocerte es el primer paso para aceptarte", "Conectar"),
                  ("Carrusel", "Cómo es tu primera sesión conmigo", "Confianza"),
                  ("Story / caja", "¿Qué llevas tiempo queriendo entender de ti?", "Conversación"),
                  ("Invitación", "Modalidades, horarios y cómo agendar", "Conversión"),
              ],
              widths=[3.5, 8.7, 4.5], font_size=8.3)


def add_scope(doc: Document) -> None:
    add_section_heading(doc, "05", "Alcance mensual",
                        "Dos niveles para ajustar el ritmo a la realidad y presupuesto de la práctica.")
    add_table(doc,
              ["Incluye", "Plan Esencial", "Plan Crecimiento · recomendado"],
              [
                  ("Piezas para feed / mes", "12 (hasta 2 Reels faceless)", "16 (hasta 4 Reels faceless)"),
                  ("Secuencias de stories / mes", "8", "12"),
                  ("Facebook + Instagram", "Diseño, texto, programación y réplica", "Diseño, texto, programación y optimización"),
                  ("Meta Ads", "Gestión de 1 objetivo de campaña", "Gestión de hasta 2 objetivos de campaña"),
                  ("Optimización", "Mensual", "Quincenal"),
                  ("Reporte", "Mensual", "Mensual + revisión estratégica"),
                  ("WhatsApp Business", "Guiones y respuestas base", "Guiones, seguimiento y ajustes"),
                  ("Configuración inicial", "Incluida en el primer mes", "Incluida en el primer mes"),
              ],
              widths=[5.1, 5.7, 5.9], font_size=8.0, first_col_bold=True)

    add_spacer(doc, 8)
    add_callout(doc, "Límite ético y operativo",
                "Aimarktech puede moderar comentarios generales y facilitar guiones, pero MaryTere atenderá los mensajes de pacientes y cualquier conversación clínica. No se solicitarán historias clínicas ni detalles sensibles por anuncios o comentarios públicos.",
                fill=CREAM, stripe=YELLOW)

    add_spacer(doc, 7)
    add_text(doc, "Lo que no está incluido", size=12, bold=True, color=INK, after=6)
    for text in [
        "Presupuesto de pauta, licencias, plataformas o servicios de terceros.",
        "Rediseño completo/vectorización del logo, sesión fotográfica o producción audiovisual externa.",
        "Sitio web, landing page, dominio, correo profesional o automatizaciones fuera del alcance descrito.",
        "Atención terapéutica, respuesta clínica o gestión de expedientes/pacientes.",
        "Asesoría jurídica, fiscal, clínica o regulatoria especializada.",
    ]:
        add_bullet(doc, text)


def add_investment(doc: Document) -> None:
    add_section_heading(doc, "06", "Inversión propuesta",
                        "Honorarios de Aimarktech y pauta publicitaria se presentan por separado para mantener total transparencia.")

    add_text(doc, "A) Honorarios de Aimarktech", size=12, bold=True, color=INK, after=6)
    add_table(doc,
              ["Plan", "Servicio mensual", "Ideal para", "Inversión"],
              [
                  ("Esencial", "Base de contenido + 1 objetivo de campaña", "Empezar con ritmo controlado", "$2,500 MXN / mes"),
                  ("Crecimiento · recomendado", "Mayor frecuencia + 2 objetivos + revisión quincenal", "Piloto de 90 días con aprendizaje más rápido", "$3,500 MXN / mes"),
              ],
              widths=[3.5, 5.5, 5.0, 2.7], font_size=8.3, first_col_bold=True)

    add_spacer(doc, 7)
    add_text(doc, "B) Pauta publicitaria (pagada directamente a Meta)", size=12, bold=True, color=INK, after=6)
    add_table(doc,
              ["Concepto", "Frecuencia", "Importe sugerido", "Forma de pago"],
              [
                  ("Facebook e Instagram Ads", "Mensual / $1,000 por quincena", "$2,000 MXN / mes", "MaryTere paga directamente a Meta"),
              ],
              widths=[5.1, 4.4, 3.5, 3.7], font_size=8.3, first_col_bold=True)

    add_spacer(doc, 8)
    add_callout(doc, "Escenario recomendado",
                "Plan Crecimiento $3,500 + pauta sugerida $2,000 = inversión mensual estimada de $5,500 MXN. La configuración inicial descrita en la sección 4 queda incluida en el primer mes del piloto; no se cobra una cuota adicional de arranque.",
                fill="EAF7ED", stripe=GREEN)

    add_spacer(doc, 7)
    add_text(doc, "Opcionales (solo con autorización previa)", size=12, bold=True, color=INK, after=6)
    add_table(doc,
              ["Opcional", "Tipo", "Precio de referencia"],
              [
                  ("Rediseño/vectorización completa del logotipo", "Único", "$1,500 MXN"),
                  ("Perfil de Empresa en Google: alta y optimización", "Único", "$1,500 MXN"),
                  ("Landing page informativa y de contacto", "Proyecto", "Desde $4,500 MXN"),
              ],
              widths=[9.0, 3.2, 4.5], font_size=8.4)

    add_spacer(doc, 7)
    add_text(doc,
             "Todos los importes están expresados en pesos mexicanos. En caso de requerir factura, se aplicarán los impuestos correspondientes. Los precios son editables y válidos durante 30 días naturales.",
             size=8.3, italic=True, color=INK_SOFT, after=4)


def add_measurement_roles(doc: Document) -> None:
    add_section_heading(doc, "07", "Medición, responsabilidades y cuidados",
                        "La estrategia se ajusta con datos; la confianza se protege con límites claros.")
    add_text(doc, "Indicadores que revisaremos", size=12, bold=True, color=INK, after=6)
    add_table(doc,
              ["Etapa", "Indicadores", "Qué nos dirán"],
              [
                  ("Reconocimiento", "Alcance, impresiones y seguidores relevantes", "Si la marca empieza a ser visible"),
                  ("Interés", "Guardados, compartidos, visitas y clics", "Qué temas generan confianza"),
                  ("Conversación", "Mensajes iniciados y costo por mensaje", "Si el llamado a WhatsApp funciona"),
                  ("Negocio", "Citas agendadas y pacientes que indican venir de redes", "Qué acciones aportan al objetivo real"),
              ],
              widths=[3.0, 6.5, 7.2], font_size=8.4, first_col_bold=True)

    add_spacer(doc, 8)
    add_text(doc, "Responsabilidades", size=12, bold=True, color=INK, after=6)
    add_table(doc,
              ["Aimarktech", "MaryTere"],
              [
                  ("Estrategia, diseño, textos, programación y gestión de campañas.", "Validar que el contenido clínico represente correctamente su práctica."),
                  ("Reporte, aprendizaje y optimización según métricas.", "Responder prospectos, confirmar disponibilidad y agendar citas."),
                  ("Protección de accesos y uso limitado a las tareas autorizadas.", "Proporcionar accesos, datos, logo y materiales con derechos de uso."),
                  ("Un ciclo de ajustes por lote de contenido.", "Aprobar o comentar cada lote dentro de 2 días hábiles."),
              ],
              widths=[8.35, 8.35], font_size=8.4)

    add_spacer(doc, 8)
    add_text(doc, "Cuidados específicos de salud mental", size=12, bold=True, color=INK, after=6)
    for text in [
        "No se garantizan pacientes, ventas ni resultados clínicos; los resultados dependen de múltiples factores.",
        "Los anuncios no señalarán atributos sensibles ni usarán frases como “tú tienes ansiedad/depresión”.",
        "Los testimonios requieren consentimiento expreso y deben minimizar datos identificables.",
        "WhatsApp se utilizará para información y agendamiento, no como sustituto de atención de crisis.",
        "Antes de captar datos sensibles deberá existir un aviso de privacidad apropiado; esta propuesta no sustituye revisión jurídica.",
        "MaryTere tendrá aprobación final de toda afirmación profesional, clínica o de credenciales.",
    ]:
        add_bullet(doc, text)


def add_terms_acceptance(doc: Document) -> None:
    add_section_heading(doc, "08", "Condiciones y siguiente paso",
                        "Un acuerdo claro antes de iniciar evita expectativas distintas durante el proyecto.")
    add_text(doc, "Condiciones comerciales", size=12, bold=True, color=INK, after=6)
    for text in [
        "El servicio se paga por mes anticipado. La pauta se paga directamente a Meta desde el método de pago de MaryTere.",
        "Se recomienda un piloto de 90 días para reunir información útil; no existe permanencia implícita. Para detener el siguiente ciclo, se solicita aviso con 15 días naturales de anticipación.",
        "Los retrasos de aprobación o entrega de accesos pueden mover el calendario sin penalización para Aimarktech.",
        "Los cambios fuera del alcance se cotizan y aprueban antes de ejecutarse.",
        "Los materiales finales pagados se entregan para uso de MaryTere; Aimarktech conserva sus metodologías y plantillas base.",
        "La información compartida se tratará de manera confidencial y solo para ejecutar el alcance autorizado.",
    ]:
        add_bullet(doc, text)

    add_spacer(doc, 7)
    add_text(doc, "Para comenzar necesitamos", size=12, bold=True, color=INK, after=6)
    for text in [
        "URL de Facebook y acceso/rol en Meta Business Suite.",
        "Número de WhatsApp dedicado y confirmación del teléfono que debe aparecer públicamente.",
        "Archivo del logo en la mejor calidad disponible.",
        "Zona aproximada del consultorio en Querétaro (sin publicar la dirección exacta).",
        "Número de cédula y forma autorizada de comunicar credenciales.",
        "Consentimiento escrito si se decide utilizar el testimonio disponible.",
        "Confirmación de precios, horarios y modalidades que se publicarán.",
    ]:
        add_bullet(doc, text)

    add_spacer(doc, 8)
    add_callout(doc, "Siguiente paso",
                "Seleccionar el plan, confirmar la pauta, realizar el primer pago y agendar una sesión de arranque de 45 minutos. Después recibiremos accesos y materiales para iniciar la Fase 0.",
                fill=BG_SOFT, stripe=BLUE)

    add_spacer(doc, 9)
    add_text(doc, "Aceptación de propuesta", size=12, bold=True, color=INK, after=6)
    selection = add_table(doc,
                          ["Selección", "Plan / concepto", "Importe"],
                          [
                              ("☐", "Plan Esencial", "$2,500 MXN / mes"),
                              ("☐", "Plan Crecimiento · recomendado", "$3,500 MXN / mes"),
                              ("☐", "Pauta publicitaria sugerida", "$2,000 MXN / mes"),
                              ("☐", "Otro / ajuste acordado: ____________________", "$________________"),
                          ],
                          widths=[2.0, 9.7, 5.0], font_size=8.6)

    add_spacer(doc, 8)
    sig = doc.add_table(rows=2, cols=3)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_layout(sig)
    widths = [7.1, 4.6, 5.0]
    for i, w in enumerate(widths):
        sig.columns[i].width = Cm(w)
    labels = ["Nombre y firma de MaryTere", "Fecha", "Fecha de inicio"]
    for i, label in enumerate(labels):
        c = sig.cell(0, i)
        c.text = ""
        set_cell_shading(c, BG_SOFT)
        set_cell_margins(c, 100, 120, 100, 120)
        p = c.paragraphs[0]
        r = p.add_run(label)
        set_run(r, size=8.0, bold=True, color=INK)
        set_cell_border(c, bottom={"val": "single", "sz": 4, "color": LINE}, left={"val": "single", "sz": 4, "color": LINE}, right={"val": "single", "sz": 4, "color": LINE}, top={"val": "single", "sz": 4, "color": LINE})
        c2 = sig.cell(1, i)
        c2.text = "\n\n"
        set_cell_margins(c2, 120, 120, 120, 120)
        set_cell_border(c2, bottom={"val": "single", "sz": 4, "color": LINE}, left={"val": "single", "sz": 4, "color": LINE}, right={"val": "single", "sz": 4, "color": LINE}, top={"val": "single", "sz": 4, "color": LINE})

    add_spacer(doc, 12)
    contact = doc.add_table(rows=1, cols=2)
    contact.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_layout(contact)
    contact.columns[0].width = Cm(10.2)
    contact.columns[1].width = Cm(6.5)
    left, right = contact.rows[0].cells
    for c in (left, right):
        set_cell_shading(c, DARK)
        set_cell_margins(c, 200, 220, 200, 220)
        set_cell_border(c, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
    left.text = ""
    p = left.paragraphs[0]
    style_paragraph(p, after=3)
    r = p.add_run("José Antonio Aguilar")
    set_run(r, size=10.5, bold=True, color=WHITE, font="Montserrat")
    p = left.add_paragraph()
    style_paragraph(p, after=2)
    r = p.add_run("Aimarktech · Sistemas digitales, IA y marketing para PyMEs")
    set_run(r, size=7.8, color="C1CEDE")
    p = left.add_paragraph()
    style_paragraph(p, after=0)
    r = p.add_run("Resultados reales, sin humo.")
    set_run(r, size=8.3, bold=True, color=YELLOW)
    right.text = ""
    for text in [PHONE, EMAIL, WEBSITE]:
        p = right.add_paragraph() if right.paragraphs[0].text else right.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style_paragraph(p, after=3)
        r = p.add_run(text)
        set_run(r, size=8.0, bold=True, color=WHITE)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "Propuesta estratégica y comercial — Psicól. MaryTere"
    props.subject = "Presencia digital y captación ética de pacientes"
    props.author = "José Antonio Aguilar · Aimarktech"
    props.keywords = "MaryTere, Psicoterapia Gestalt, Aimarktech, redes sociales, Meta Ads"
    props.comments = "Documento editable. Versión 1.0 · Agosto 2026."

    add_cover(doc)
    add_executive_summary(doc)
    doc.add_page_break()
    add_diagnosis(doc)
    doc.add_page_break()
    add_strategy(doc)
    doc.add_page_break()
    add_roadmap(doc)
    doc.add_page_break()
    add_scope(doc)
    doc.add_page_break()
    add_investment(doc)
    doc.add_page_break()
    add_measurement_roles(doc)
    doc.add_page_break()
    add_terms_acceptance(doc)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
